"""Turn a Word upload into a PDF for the AI review.

Prefers PyMuPDF's DOCX converter at the document's real page size. If the host
build cannot open Word files (or returns the tiny 400x600 default canvas), fall
back to a python-docx render that keeps Word point sizes (11pt stays 11pt).
"""

from __future__ import annotations

import io
import re
from typing import Any

import fitz

EMU_PER_INCH = 914400.0
PT_PER_INCH = 72.0


def docx_to_pdf(raw: bytes) -> bytes:
    width, height = page_points(raw)
    try:
        pdf_bytes = _fitz_convert(raw, width, height)
        probe = fitz.open(stream=pdf_bytes, filetype="pdf")
        try:
            page = probe[0]
            # PyMuPDF's default DOCX canvas is 400x600. That exaggerates fonts
            # when pasted into a Letter review page — reject it.
            if page.rect.width < width * 0.85 or page.rect.height < height * 0.85:
                raise RuntimeError("DOCX converter used a too-small page canvas")
        finally:
            probe.close()
        return pdf_bytes
    except Exception:
        return render_docx_pdf(raw)


def page_points(raw: bytes) -> tuple[float, float]:
    try:
        from docx import Document

        section = Document(io.BytesIO(raw)).sections[0]
        width = float(section.page_width) / EMU_PER_INCH * PT_PER_INCH
        height = float(section.page_height) / EMU_PER_INCH * PT_PER_INCH
        if width >= 200 and height >= 200:
            return width, height
    except Exception:
        pass
    return 612.0, 792.0


def _fitz_convert(raw: bytes, width: float, height: float) -> bytes:
    attempts = (
        {"filetype": "docx", "rect": fitz.Rect(0, 0, width, height)},
        {"filetype": "docx", "width": width, "height": height},
        {"filetype": "docx"},
    )
    last_error: Exception | None = None
    for kwargs in attempts:
        try:
            document = fitz.open(stream=raw, **kwargs)
            try:
                return document.convert_to_pdf()
            finally:
                document.close()
        except Exception as exc:
            last_error = exc
    raise RuntimeError(str(last_error or "DOCX conversion failed"))


def render_docx_pdf(raw: bytes) -> bytes:
    """Layout-preserving fallback using python-docx run sizes and page margins."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    document = Document(io.BytesIO(raw))
    section = document.sections[0]
    width = float(section.page_width) / EMU_PER_INCH * PT_PER_INCH
    height = float(section.page_height) / EMU_PER_INCH * PT_PER_INCH
    margin_left = float(section.left_margin) / EMU_PER_INCH * PT_PER_INCH
    margin_right = float(section.right_margin) / EMU_PER_INCH * PT_PER_INCH
    margin_top = float(section.top_margin) / EMU_PER_INCH * PT_PER_INCH
    margin_bottom = float(section.bottom_margin) / EMU_PER_INCH * PT_PER_INCH

    left = max(36.0, margin_left)
    right = width - max(36.0, margin_right)
    top = max(36.0, margin_top)
    bottom = height - max(36.0, margin_bottom)
    content_w = max(120.0, right - left)

    output = fitz.open()
    page = output.new_page(width=width, height=height)
    y = top
    default_size = 11.0

    for paragraph in document.paragraphs:
        text = paragraph.text or ""
        if not text.strip() and not paragraph.runs:
            y += default_size * 0.8
            if y > bottom:
                page = output.new_page(width=width, height=height)
                y = top
            continue

        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        base_size = _style_size(style_name, default_size)
        align = paragraph.alignment
        runs = list(paragraph.runs) or [None]
        pieces: list[tuple[str, float, bool, bool]] = []
        for run in runs:
            if run is None:
                pieces.append((text, base_size, False, False))
                continue
            chunk = run.text or ""
            if not chunk:
                continue
            size = _run_size(run, base_size)
            pieces.append((chunk, size, bool(run.bold), bool(run.italic)))
        if not pieces:
            continue

        lines = _wrap_runs(pieces, content_w)
        for line in lines:
            line_height = max((size for _t, size, _b, _i in line), default=base_size) * 1.25
            if y + line_height > bottom:
                page = output.new_page(width=width, height=height)
                y = top
            line_width = sum(_text_width(t, size, bold, italic) for t, size, bold, italic in line)
            if align == WD_ALIGN_PARAGRAPH.CENTER:
                x = left + max(0.0, (content_w - line_width) / 2)
            elif align == WD_ALIGN_PARAGRAPH.RIGHT:
                x = left + max(0.0, content_w - line_width)
            else:
                x = left
            baseline = y + line_height * 0.8
            for chunk, size, bold, italic in line:
                fontname = _font_name(bold, italic)
                page.insert_text(
                    (x, baseline),
                    chunk,
                    fontsize=size,
                    fontname=fontname,
                    color=(0.1, 0.1, 0.1),
                )
                x += _text_width(chunk, size, bold, italic)
            y += line_height
        y += base_size * 0.35

    # Tables: simple sequential text so content is not dropped.
    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
            line = " | ".join(part for part in cells if part)
            if not line:
                continue
            size = default_size
            wrapped = _wrap_plain(line, content_w, size)
            for part in wrapped:
                line_height = size * 1.25
                if y + line_height > bottom:
                    page = output.new_page(width=width, height=height)
                    y = top
                page.insert_text((left, y + line_height * 0.8), part, fontsize=size, fontname="helv", color=(0.1, 0.1, 0.1))
                y += line_height
            y += size * 0.3

    if output.page_count == 0:
        output.new_page(width=width, height=height)

    data = output.tobytes()
    output.close()
    return data


def _style_size(style_name: str, default: float) -> float:
    name = style_name.lower()
    if "heading 1" in name or name == "title":
        return 20.0
    if "heading 2" in name:
        return 16.0
    if "heading 3" in name:
        return 14.0
    if "heading" in name:
        return 13.0
    return default


def _run_size(run: Any, default: float) -> float:
    from docx.oxml.ns import qn

    try:
        if run.font is not None and run.font.size is not None:
            return float(run.font.size.pt)
    except Exception:
        pass
    try:
        style = run.style
        if style is not None and style.font is not None and style.font.size is not None:
            return float(style.font.size.pt)
    except Exception:
        pass
    try:
        element = run._element
        rpr = element.find(qn("w:rPr")) if element is not None else None
        if rpr is not None:
            sz = rpr.find(qn("w:sz"))
            if sz is not None and sz.get(qn("w:val")):
                return float(sz.get(qn("w:val"))) / 2.0
    except Exception:
        pass
    return default


def _font_name(bold: bool, italic: bool) -> str:
    if bold and italic:
        return "hebi"
    if bold:
        return "hebo"
    if italic:
        return "heit"
    return "helv"


def _text_width(text: str, size: float, bold: bool, italic: bool) -> float:
    font = fitz.Font(_font_name(bold, italic))
    return font.text_length(text, fontsize=size)


def _wrap_runs(pieces: list[tuple[str, float, bool, bool]], width: float) -> list[list[tuple[str, float, bool, bool]]]:
    lines: list[list[tuple[str, float, bool, bool]]] = [[]]
    line_width = 0.0
    for text, size, bold, italic in pieces:
        parts = re.split(r"(\s+)", text)
        for part in parts:
            if not part:
                continue
            part_w = _text_width(part, size, bold, italic)
            if lines[-1] and line_width + part_w > width and not part.isspace():
                lines.append([])
                line_width = 0.0
                if part.isspace():
                    continue
            lines[-1].append((part, size, bold, italic))
            line_width += part_w
    return [line for line in lines if line]


def _wrap_plain(text: str, width: float, size: float) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        trial = word if not current else f"{current} {word}"
        if _text_width(trial, size, False, False) <= width:
            current = trial
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]
